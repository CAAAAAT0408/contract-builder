"""변경내역 엑셀 만들기."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

KIND = {'update': '변경', 'add': '추가', 'contract_only': '계약서 전용 문구', 'extra': '계약서 전용 괄호문구',
        'dong_replace': '동별 교체', 'info': '참고'}


def _text(segs, mark):
    out = []
    for t, h in segs or []:
        out.append(f'[{t}]' if (h and mark) else t)
    return ''.join(out)


def make_report(plan, decisions, fields, rules, log, names):
    wb = Workbook()
    ws = wb.active
    ws.title = '제20조 변경내역'
    head = ['번호', '구분', '종류', '적용', '변경 전', '변경 후 ([ ]=바뀐 글자)', '비고']
    ws.append(head)
    for ch in plan['changes']:
        on = decisions.get(ch['id'], ch['default'])
        kind = ch['kind']
        if kind == 'contract_only':
            applied = '삭제' if on else '유지'
        elif kind == 'extra':
            applied = '괄호문구 삭제' if on else '유지'
        elif kind == 'info':
            applied = '-'
        else:
            applied = '적용' if on else '제외'
        if kind == 'dong_replace':
            after = '\n'.join(f"{b['dong']}\n" + '\n'.join(b['items']) for b in ch['blocks'])
        else:
            after = _text(ch.get('after'), True)
        ws.append([ch['id'][1:], ch.get('group', ''), KIND.get(kind, kind), applied,
                   _text(ch.get('before'), False), after, (ch.get('title') or '') + (' / ' + ch['note'] if ch.get('note') else '')])
    widths = [6, 18, 14, 10, 70, 70, 40]
    for i, wd in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    fill = PatternFill('solid', fgColor='1F4E79')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill
    ws.freeze_panes = 'A2'

    ws2 = wb.create_sheet('1페이지·이름')
    ws2.append(['항목', '샘플 계약서', '새 값', '적용'])
    for f in plan['fields']:
        v = fields.get(f['key'])
        ws2.append([f['label'], f['old'], v if v is not None else f['new'], '적용' if v is not None else '제외'])
    for r in rules:
        ws2.append([f"이름 바꾸기 ({r.get('count', '')}곳)", r.get('old', ''), r.get('new', ''),
                    '적용' if r.get('apply') else '제외'])
    ws2.append([])
    ws2.append(['샘플 계약서', names.get('contract', '')])
    ws2.append(['모집공고', names.get('notice', '')])
    for col, wd in zip('ABCD', (28, 40, 40, 8)):
        ws2.column_dimensions[col].width = wd
    for c in ws2[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill

    if plan.get('warnings'):
        ws3 = wb.create_sheet('확인할 점')
        for wmsg in plan['warnings']:
            ws3.append([wmsg])
        ws3.column_dimensions['A'].width = 140
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
